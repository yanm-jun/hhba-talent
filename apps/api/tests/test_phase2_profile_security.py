import os, uuid
RUN=uuid.uuid4().hex
os.environ['DATABASE_URL']=f'sqlite:///./phase2-security-{RUN}.sqlite3'
os.environ['HHBA_BOOTSTRAP_ADMIN_EMAIL']=f'admin-{RUN}@hhba-synthetic.com'
from fastapi.testclient import TestClient
from app.main import app, Base, engine
Base.metadata.create_all(engine)
def login(c,email):
 r=c.post('/api/v1/auth/session',json={'email':email,'password':'SyntheticOnly123!'});assert r.status_code==200;return r.json()['csrf_token']
def test_profile_publication_and_withdrawal_are_server_enforced():
 a=TestClient(app);b=TestClient(app);ea=f'a-{RUN}@hhba-synthetic.com';eb=f'b-{RUN}@hhba-synthetic.com'
 for c,e in [(a,ea),(b,eb)]:assert c.post('/api/v1/auth/register',json={'email':e,'password':'SyntheticOnly123!','display_name':'模拟'}).status_code==201
 ca=login(a,ea);cb=login(b,eb)
 assert a.put('/api/v1/me/profile',json={'display_name':'合成用户','content':{'headline':'测试','skills':['Python'],'services':['咨询']}},headers={'X-CSRF-Token':ca}).status_code==200
 assert a.post('/api/v1/me/profile/publish',headers={'X-CSRF-Token':ca}).status_code==403
 assert a.post('/api/v1/me/profile/confirm',headers={'X-CSRF-Token':ca}).status_code==200
 assert a.put('/api/v1/me/consents/public_share',json={'granted':True},headers={'X-CSRF-Token':ca}).status_code==200
 p=a.post('/api/v1/me/profile/publish',headers={'X-CSRF-Token':ca});assert p.status_code==200;slug=p.json()['slug']
 public=a.get('/api/v1/public/profiles/'+slug);assert public.status_code==200 and '联系方式' not in public.text
 assert b.get('/api/v1/me/resumes/'+str(uuid.uuid4())).status_code==404
 assert a.put('/api/v1/me/consents/public_share',json={'granted':False},headers={'X-CSRF-Token':ca}).status_code==200
 assert a.get('/api/v1/public/profiles/'+slug).status_code==404

def test_private_upload_is_owner_scoped_and_signature_checked():
 a=TestClient(app);b=TestClient(app);ea=f'upload-a-{RUN}@hhba-synthetic.com';eb=f'upload-b-{RUN}@hhba-synthetic.com'
 for c,e in [(a,ea),(b,eb)]:assert c.post('/api/v1/auth/register',json={'email':e,'password':'SyntheticOnly123!','display_name':'模拟'}).status_code==201
 ca=login(a,ea);login(b,eb)
 bad=a.post('/api/v1/me/resumes',files={'file':('bad.pdf',b'not a pdf','application/pdf')},headers={'X-CSRF-Token':ca});assert bad.status_code==422
 up=a.post('/api/v1/me/resumes',files={'file':('synthetic.pdf',b'%PDF-1.4\n% synthetic only','application/pdf')},headers={'X-CSRF-Token':ca});assert up.status_code==201;rid=up.json()['resume_id']
 assert b.get('/api/v1/me/resumes/'+rid).status_code==404
 assert a.delete('/api/v1/me/resumes/'+rid,headers={'X-CSRF-Token':ca}).status_code==200
 assert a.get('/api/v1/me/resumes/'+rid).status_code==404

def test_synthetic_docx_upload_and_local_parse_state():
 from docx import Document
 import io
 doc=Document();doc.add_heading('Synthetic Resume');doc.add_paragraph('Python\nReact\nSQL');buf=io.BytesIO();doc.save(buf)
 c=TestClient(app);email=f'docx-{RUN}@hhba-synthetic.com'
 assert c.post('/api/v1/auth/register',json={'email':email,'password':'SyntheticOnly123!','display_name':'模拟'}).status_code==201
 csrf=login(c,email);up=c.post('/api/v1/me/resumes',files={'file':('synthetic.docx',buf.getvalue(),'application/vnd.openxmlformats-officedocument.wordprocessingml.document')},headers={'X-CSRF-Token':csrf});assert up.status_code==201
 rid=up.json()['resume_id'];assert c.get('/api/v1/me/resumes/'+rid+'/download').status_code==200
 parsed=c.post('/api/v1/me/resumes/'+rid+'/parse',headers={'X-CSRF-Token':csrf});assert parsed.status_code==200
 assert parsed.json()['state']=='queued', parsed.json()
 from app.docling_worker import run_once
 assert run_once() is True
 status=c.get('/api/v1/me/resumes/'+rid);assert status.json()['job_state']=='parsed',status.json()
 applied=c.post('/api/v1/me/parsing-jobs/'+up.json()['job_id']+'/apply-draft',headers={'X-CSRF-Token':csrf});assert applied.status_code==200
 assert c.get('/api/v1/me/profile').json()['profile']['status']=='draft'

def test_synthetic_pdf_upload_and_local_parse_state():
 from reportlab.pdfgen.canvas import Canvas
 import io
 buf=io.BytesIO();pdf=Canvas(buf);pdf.drawString(72,720,'Synthetic Resume - Python React SQL');pdf.save()
 c=TestClient(app);email=f'pdf-{RUN}@hhba-synthetic.com'
 assert c.post('/api/v1/auth/register',json={'email':email,'password':'SyntheticOnly123!','display_name':'模拟'}).status_code==201
 csrf=login(c,email);up=c.post('/api/v1/me/resumes',files={'file':('synthetic.pdf',buf.getvalue(),'application/pdf')},headers={'X-CSRF-Token':csrf});assert up.status_code==201
 parsed=c.post('/api/v1/me/resumes/'+up.json()['resume_id']+'/parse',headers={'X-CSRF-Token':csrf});assert parsed.status_code==200
 assert parsed.json()['state']=='queued',parsed.json()
 from app.docling_worker import run_once
 assert run_once() is True
 assert c.get('/api/v1/me/resumes/'+up.json()['resume_id']).json()['job_state']=='parsed'

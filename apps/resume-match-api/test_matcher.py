"""No real resume or DeepSeek token is needed for these safety tests."""
import io
import json
import os
import unittest
from unittest.mock import AsyncMock, patch

from fastapi.testclient import TestClient
from docx import Document

import matcher
import server


class SourceTests(unittest.TestCase):
    def setUp(self):
        self.resume = ("本科 信息管理专业。3年销售经验，企业销售，负责客户拜访及 Salesforce。"
                       "了解 Excel 数据分析，目标职位是 Sales Account Executive。")
        self.profile = {
            "education": "BACHELOR", "experience_years": 3,
            "skills": [{"name": "excel", "evidence": "Excel"}, {"name": "salesforce", "evidence": "Salesforce"}],
            "target_roles": ["Sales Account Executive"],
        }
        self.job = {
            "id": "greenhouse-stripe-123456",
            "platform": "greenhouse", "source": "stripe", "company": "Stripe",
            "title": "Enterprise Account Executive", "location": "Singapore",
            "url": "https://stripe.com/jobs/search?gh_jid=123456", "remote": False,
            "domestic": False,
        }

    def test_profile_rejects_llm_hallucinations(self):
        proposed = {
            "education": {"level": "PHD", "evidence": "PhD Stanford"},
            "experience_years": 20, "experience_evidence": "20 years of experience",
            "skills": [
                {"name": "Excel", "evidence": "Excel"},
                {"name": "Kubernetes", "evidence": "Kubernetes"},
                {"name": "Salesforce", "evidence": "Salesforce"},
            ],
            "target_roles": [
                {"role": "Account Executive", "evidence": "Sales Account Executive"},
                {"role": "Chief Data Scientist", "evidence": "Chief Data Scientist"},
            ],
        }
        profile = matcher.profile_from_llm(proposed, self.resume)
        self.assertIsNone(profile["education"])
        self.assertIsNone(profile["experience_years"])
        self.assertEqual([x["name"] for x in profile["skills"]], ["excel", "salesforce"])
        self.assertEqual(profile["target_roles"], ["Account Executive"])

    def test_url_validation_no_external_redirect_or_guessed_id(self):
        self.assertTrue(matcher.valid_official_url(self.job))
        self.assertFalse(matcher.valid_official_url({**self.job, "url": "https://evil.com/jobs/search?gh_jid=123456"}))
        self.assertFalse(matcher.valid_official_url({**self.job, "url": "https://stripe.com/jobs/search?gh_jid=987654"}))
        self.assertFalse(matcher.valid_official_url({**self.job, "url": "http://stripe.com/jobs/search?gh_jid=123456"}))
        self.assertFalse(matcher.valid_official_url({**self.job, "url": "https://stripe.com@evil.com/jobs/search?gh_jid=123456"}))
        self.assertTrue(matcher.valid_official_url({
            "id": "ashby-deepgram-abc123", "source": "deepgram", "platform": "ashby",
            "url": "https://jobs.ashbyhq.com/deepgram/abc123",
        }))

    def test_matching_grounded_by_live_audit_and_explicit_requirements(self):
        facts = {
            self.job["id"]: {"countryCode": "SG", "original_source_hash": "a" * 64, "education": "BACHELOR",
                             "experienceMinYears": 2, "skills": ["excel", "salesforce"]},
        }
        audits = {self.job["id"]: {"status": "OFFICIAL_JD_AND_REQUIREMENTS_FOUND", "full_original_jd_verified": True, "jd_sha256": "a" * 64}}
        result = matcher.rank_official_jobs(self.profile, [self.job], facts, audits, country="SG")
        self.assertEqual(len(result), 1)
        self.assertEqual(result[0]["official_url"], self.job["url"])
        self.assertTrue(result[0]["matched"])
        self.assertEqual(result[0]["not_evidenced"], [])
        self.assertIn("needs_verification", result[0])
        self.assertLessEqual(result[0]["reference_score"], 95)
        self.assertEqual(matcher.rank_official_jobs(self.profile, [self.job], facts, audits, country="US"), [])
        self.assertEqual(matcher.rank_official_jobs(
            self.profile, [self.job], facts, {self.job["id"]: {"status": "MISSING_FROM_OFFICIAL_BOARD"}}), [])
        self.assertEqual(matcher.rank_official_jobs(
            self.profile, [self.job], {}, audits, remote_only=True), [])

    def test_unknown_is_never_assumed_degree_compliant(self):
        f = {self.job["id"]: {"countryCode": "SG", "original_source_hash": "a" * 64, "education": None,
                              "experienceMinYears": None, "skills": []}}
        result = matcher.rank_official_jobs(self.profile, [self.job], f,
                         {self.job["id"]: {"status": "OFFICIAL_JD_AND_REQUIREMENTS_FOUND", "full_original_jd_verified": True, "jd_sha256": "a" * 64}})
        self.assertEqual(len(result), 1)
        self.assertTrue(any("不能当作不限学历" in s for s in result[0]["needs_verification"]))

    def test_no_unsupported_platform(self):
        j = dict(self.job, platform="notanats", id="notanats-stripe-123456")
        self.assertFalse(matcher.valid_official_url(j))

    def test_pdf_and_docx_types_and_pii(self):
        d = Document()
        d.add_paragraph("本科销售顾问。管理多个客户项目以及业务增长。Excel 和 Salesforce，三年工作经验。工作内容包括市场销售、客户关系管理、门店培训、数据处理、项目交付、业务沟通和团队合作。")
        output = io.BytesIO()
        d.save(output)
        txt = server.parse_resume(output.getvalue(), "resume.docx")
        self.assertIn("Salesforce", txt)
        masked = server.redact_resume("李先生 13800138000 email@example.com https://github.com/some-user")
        self.assertNotIn("13800138000", masked)
        self.assertNotIn("email@example.com", masked)
        self.assertNotIn("github.com", masked)

    def test_requires_admin_key_even_before_file_read_and_privacy_consents(self):
        client = TestClient(server.app)
        with patch.object(server, "MATCH_PUBLIC_ENABLED", True), patch.dict(os.environ, {"DEEPSEEK_API_KEY": ""}):
            reply = client.post("/api/v1/match/resume", files={"file": (
                "cv.docx", b"PK\x03\x04", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )}, data={"consent": "true"}, headers={"Origin": "https://yanm-jun.github.io"})
            self.assertEqual(reply.status_code, 503)
            self.assertEqual(reply.json()["error"], "ADMIN_AI_NOT_CONFIGURED")
        with patch.object(server, "MATCH_PUBLIC_ENABLED", True), patch.dict(os.environ, {"DEEPSEEK_API_KEY": "test-not-real"}):
            reply = client.post("/api/v1/match/resume", files={"file": (
                "cv.pdf", b"%PDF-notreal", "application/pdf"
            )}, data={"consent": "false"}, headers={"Origin": "https://yanm-jun.github.io"})
            self.assertEqual(reply.status_code, 400)
            self.assertEqual(reply.json()["error"], "EXPLICIT_AI_SHARING_CONSENT_REQUIRED")
            spoofed = client.post("/api/v1/match/resume", files={"file": (
                "cv.pdf", b"%PDF-notreal", "application/pdf"
            )}, data={"consent": "true"}, headers={"Origin": "https://evil.com"})
            self.assertEqual(spoofed.status_code, 403)

    def test_reject_wrong_magic_or_scanned_pdfs(self):
        with self.assertRaises(Exception):
            server.parse_resume(b"evil not pdf" * 20, "cv.pdf")
        with self.assertRaises(Exception):
            server.parse_resume(b"PK\x03\x04broken" * 20, "cv.docx")

    def test_complete_upload_to_official_match_with_stubbed_platform_deepseek(self):
        """Integration test: users send no key; only server knows the configured key."""
        document = Document()
        document.add_paragraph(
            "Account Executive, enterprise sales and customer relationship management. "
            "I have 3 years of sales experience with Salesforce, Excel, team training and "
            "quarterly revenue targets. My role is Sales Account Executive. "
            "Contact: private@example.com."
        )
        out = io.BytesIO()
        document.save(out)
        official_hash = "a" * 64
        source = {**self.job}
        source_fact = {
            "countryCode": "SG", "cityKey": "singapore", "education": None,
            "experienceMinYears": 2, "skills": ["excel", "salesforce"],
            "original_source_hash": official_hash,
        }
        source_audit = {
            "status": "OFFICIAL_JD_AND_REQUIREMENTS_FOUND",
            "full_original_jd_verified": True, "jd_sha256": official_hash,
        }
        ai_result = {
            "education": None, "education_evidence": None,
            "experience_years": 3, "experience_evidence": "3 years of sales experience",
            "skills": [{"name": "excel", "evidence": "Excel"},
                       {"name": "salesforce", "evidence": "Salesforce"}],
            "target_roles": ["Sales Account Executive"], "warnings": [],
        }
        with (patch.dict(os.environ, {"DEEPSEEK_API_KEY": "test-server-only-key"}),
              patch.object(server, "MATCH_PUBLIC_ENABLED", True),
              patch.object(server, "reserve_budget", new_callable=AsyncMock),
              patch.object(server, "release_slot", new_callable=AsyncMock),
              patch.object(server, "download_official_data",
                           new_callable=AsyncMock, return_value=(
                               [source], {source["id"]: source_fact}, {source["id"]: source_audit})),
              patch.object(server, "extract_profile_with_deepseek",
                           new_callable=AsyncMock, return_value=ai_result) as provider):
            response = TestClient(server.app).post(
                "/api/v1/match/resume",
                headers={"Origin": "https://yanm-jun.github.io"},
                files={"file": ("resume.docx", out.getvalue(),
                                "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
                data={"consent": "true", "country": "SG", "role_hint": "sales", "limit": "10"}
            )
        self.assertEqual(response.status_code, 200, response.text)
        body = response.json()
        self.assertEqual(body["status"], "COMPLETE")
        self.assertEqual(len(body["matches"]), 1)
        self.assertEqual(body["matches"][0]["official_url"], source["url"])
        self.assertTrue(body["matches"][0]["matched"])
        self.assertEqual(body["matches"][0]["source"], "stripe")
        self.assertNotIn("test-server-only-key", response.text)
        self.assertNotIn("private@example.com", provider.call_args.args[0])
        self.assertIn("[邮箱已隐藏]", provider.call_args.args[0])

    def test_health_does_not_return_admin_secret(self):
        with patch.dict(os.environ, {"DEEPSEEK_API_KEY": "a-test-private-key"}):
            resp = TestClient(server.app).get("/health")
        self.assertTrue(resp.json()["admin_deepseek_configured"])
        self.assertNotIn("a-test-private-key", resp.text)


if __name__ == "__main__":
    unittest.main()
